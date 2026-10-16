from importlib import import_module
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.apps import apps
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase, APIClient

from .models import Account, FileConfig, Project, RemoteFile
from .projects import save_project
from .project_transfer import IMPORT_COLUMNS, validate_rows


class ProjectManagementTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username='project-user')
        self.client.force_authenticate(self.user)
        self.account = Account.objects.create(name='SSH', username='root')
        self.config = FileConfig.objects.create(
            project='old', host_name='node', ip_address='192.0.2.10',
            machine_name='machine', file_path='/srv/files', account=self.account,
        )
        self.file = RemoteFile.objects.create(
            file_config=self.config, project='old', name='x', relative_path='x',
        )

    def test_migration_backfills_legacy_projects(self):
        FileConfig.objects.filter(pk=self.config.pk).update(project_node=None)
        import_module('dashboard.migrations.0006_populate_projects').populate_projects(apps, None)
        self.config.refresh_from_db()
        self.assertEqual(self.config.project_node.name, 'old')

    def test_nested_projects_rename_and_cached_file_labels(self):
        parent = save_project({'name': 'parent'})
        child = save_project({'name': 'child', 'parent_id': parent.pk})
        response = self.client.patch(
            f'/api/file-configs/{self.config.pk}/',
            {'project_id': child.pk, 'revision': self.config.revision}, format='json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['project_path'], 'parent/child')
        renamed = self.client.patch(f'/api/projects/{child.pk}/', {'name': 'new-child'}, format='json')
        self.assertEqual(renamed.status_code, 200)
        self.config.refresh_from_db()
        self.file.refresh_from_db()
        self.assertEqual((self.config.project, self.file.project), ('new-child', 'new-child'))
        self.assertEqual(self.config.project_node.path, 'parent/new-child')

    def test_rejects_cycle_duplicate_and_nonempty_deletion(self):
        root = self.config.project_node
        child = save_project({'name': 'child', 'parent_id': root.pk})
        self.assertEqual(self.client.patch(f'/api/projects/{root.pk}/', {'parent_id': child.pk}).status_code, 400)
        self.assertEqual(self.client.post('/api/projects/', {'name': 'child', 'parent_id': root.pk}).status_code, 400)
        self.assertEqual(self.client.delete(f'/api/projects/{root.pk}/').status_code, 400)
        self.assertEqual(self.client.delete(f'/api/projects/{child.pk}/').status_code, 204)

    def test_batch_operations_check_each_revision_and_delete_cache(self):
        target = save_project({'name': 'target'})
        stale = self.client.post('/api/file-configs/batch/', {
            'action': 'move', 'project_id': target.pk,
            'items': [{'id': self.config.pk, 'revision': 88}],
        }, format='json')
        self.assertFalse(stale.data['results'][0]['success'])
        moved = self.client.post('/api/file-configs/batch/', {
            'action': 'move', 'project_id': target.pk,
            'items': [{'id': self.config.pk, 'revision': self.config.revision}],
        }, format='json')
        self.assertTrue(moved.data['results'][0]['success'])
        self.config.refresh_from_db()
        self.file.refresh_from_db()
        self.assertEqual(self.file.project, 'target')
        deleted = self.client.post('/api/file-configs/batch/', {
            'action': 'delete',
            'items': [{'id': self.config.pk, 'revision': self.config.revision}],
        }, format='json')
        self.assertTrue(deleted.data['results'][0]['success'])
        self.assertFalse(RemoteFile.objects.filter(pk=self.file.pk).exists())

    @patch('dashboard.project_views.inspect_machine')
    @patch('dashboard.project_views.connect')
    def test_inspection_collects_specs_without_touching_file_cache(self, connect, inspect):
        inspect.return_value = {
            'machine_name': 'detected', 'platform_type': 'linux', 'cpu': 4,
            'memory_bytes': 8 * 1024 ** 3, 'disk_bytes': 100 * 1024 ** 3,
            'system': 'Debian', 'arch': 'x86_64',
        }
        result = self.client.post(f'/api/file-configs/{self.config.pk}/inspect/')
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.data['config']['specifications']['cpu'], 4)
        self.assertTrue(RemoteFile.objects.filter(pk=self.file.pk).exists())
        connect.return_value.close.assert_called_once()

    def test_new_project_host_needs_no_machine_name_or_file_path(self):
        project = save_project({'name': 'new-project'})
        response = self.client.post('/api/file-configs/', {
            'project_id': project.pk, 'host_name': 'new-node', 'ip_address': '192.0.2.20',
            'port': 22, 'account_id': self.account.pk,
        }, format='json')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['machine_name'], '')
        self.assertEqual(response.data['file_path'], '')
        config = FileConfig.objects.get(pk=response.data['id'])
        self.assertEqual(config.machine_name, '')
        self.assertEqual(config.file_path, '')

    def test_host_edit_preserves_collected_name_and_legacy_file_path(self):
        response = self.client.patch(f'/api/file-configs/{self.config.pk}/', {
            'host_name': 'renamed-node', 'revision': self.config.revision,
        }, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['machine_name'], 'machine')
        self.assertEqual(response.data['file_path'], '/srv/files')
        self.assertTrue(RemoteFile.objects.filter(pk=self.file.pk).exists())

    def test_connection_change_resets_collected_name_and_specs(self):
        self.config.specifications = {'cpu': 4}
        self.config.connection_status = 'success'
        self.config.save()
        response = self.client.patch(f'/api/file-configs/{self.config.pk}/', {
            'ip_address': '192.0.2.21', 'revision': self.config.revision,
        }, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['machine_name'], '')
        self.assertEqual(response.data['specifications'], {})
        self.assertEqual(response.data['connection_status'], 'pending')

    @patch('dashboard.project_views.inspect_machine')
    @patch('dashboard.project_views.connect')
    def test_inspection_populates_machine_name_from_hostname(self, connect, inspect):
        self.config.machine_name = ''
        self.config.save(update_fields=['machine_name'])
        inspect.return_value = {
            'machine_name': 'remote-host', 'platform_type': 'linux', 'cpu': 4,
            'memory_bytes': 8 * 1024 ** 3, 'disk_bytes': 100 * 1024 ** 3,
            'system': 'Debian', 'arch': 'x86_64',
        }
        response = self.client.post(f'/api/file-configs/{self.config.pk}/inspect/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['config']['machine_name'], 'remote-host')
        self.config.refresh_from_db()
        self.assertEqual(self.config.machine_name, 'remote-host')

    @patch('dashboard.views.connect')
    def test_refresh_without_file_path_is_rejected_without_ssh_or_cache_changes(self, connect):
        self.config.file_path = ''
        self.config.save(update_fields=['file_path'])
        cached = RemoteFile.objects.create(
            file_config=self.config, project='old', name='cached', relative_path='cached',
        )
        response = self.client.post(f'/api/file-configs/{self.config.pk}/refresh/')
        self.assertEqual(response.status_code, 400)
        self.assertIn('文件路径', response.data['detail'])
        connect.assert_not_called()
        self.assertTrue(RemoteFile.objects.filter(pk=cached.pk).exists())

    @patch('dashboard.project_views.inspect_machine', side_effect=ValueError('lsblk unavailable'))
    @patch('dashboard.project_views.connect')
    def test_partial_inspection_preserves_specs_and_verifies_ssh(self, connect, inspect):
        self.config.specifications = {'cpu': 8}
        self.config.save(update_fields=['specifications'])
        result = self.client.post(f'/api/file-configs/{self.config.pk}/inspect/')
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.data['config']['connection_status'], 'success')
        self.assertEqual(result.data['config']['specifications'], {'cpu': 8})
        self.assertEqual(result.data['config']['inspection_status'], 'error')

    @patch('dashboard.project_views.inspect_machine')
    @patch('dashboard.project_views.connect')
    def test_old_inspection_cannot_overwrite_new_host_edit(self, connect, inspect):
        def edit_while_inspecting(_client):
            FileConfig.objects.filter(pk=self.config.pk).update(
                revision=self.config.revision + 2, machine_name='newer',
            )
            return {'machine_name': 'stale', 'platform_type': 'linux', 'cpu': 2}

        inspect.side_effect = edit_while_inspecting
        result = self.client.post(f'/api/file-configs/{self.config.pk}/inspect/')
        self.assertEqual(result.status_code, 409)
        self.config.refresh_from_db()
        self.assertEqual(self.config.machine_name, 'newer')

    @patch('dashboard.project_views.connect', side_effect=ConnectionError('authentication failed'))
    def test_failed_connection_preserves_cached_files(self, connect):
        result = self.client.post(f'/api/file-configs/{self.config.pk}/inspect/')
        self.assertEqual(result.status_code, 400)
        self.assertEqual(result.data['config']['connection_status'], 'error')
        self.assertTrue(RemoteFile.objects.filter(pk=self.file.pk).exists())

    def test_excel_preview_confirm_and_safe_export(self):
        book = Workbook()
        sheet = book.active
        sheet.append([label for label, _ in IMPORT_COLUMNS])
        sheet.append(['parent/child', 'new', '192.0.2.11', 22, 'SSH'])
        output = BytesIO()
        book.save(output)
        preview = self.client.post('/api/file-configs/import/', {
            'file': SimpleUploadedFile('hosts.xlsx', output.getvalue()),
        }, format='multipart')
        self.assertEqual(preview.status_code, 200)
        self.assertTrue(preview.data['valid'])
        result = self.client.post('/api/file-configs/import/', {'token': preview.data['token']}, format='json')
        self.assertEqual(result.data['created'], 1)
        imported = FileConfig.objects.get(host_name='new')
        self.assertEqual(imported.project_node.path, 'parent/child')
        self.assertEqual((imported.machine_name, imported.file_path), ('', ''))
        exported = self.client.post('/api/file-configs/export/', {'ids': [imported.pk]}, format='json')
        self.assertEqual(exported.status_code, 200)
        self.assertNotIn(b'encrypted_password', exported.content)
        sheet = load_workbook(BytesIO(exported.content), read_only=True).active
        self.assertEqual(sheet['A2'].value, 'parent/child')
        self.assertEqual(sheet['F1'].value, '机器名称')
        self.assertNotIn('文件路径', [cell.value for cell in sheet[1]])

        template = self.client.get('/api/file-configs/import-template/')
        template_sheet = load_workbook(BytesIO(template.content), read_only=True).active
        headers = [cell.value for cell in template_sheet[1]]
        self.assertNotIn('机器名称', headers)
        self.assertNotIn('文件路径', headers)

    def test_invalid_excel_row_creates_nothing(self):
        rows = [{
            'project_path': 'a/b', 'host_name': 'bad', 'ip_address': 'invalid',
            'machine_name': 'm', 'port': '22', 'account_name': 'SSH',
            'file_path': '/srv', 'row': 2,
        }]
        self.assertEqual(validate_rows(rows)[0]['status'], 'error')
        self.assertFalse(Project.objects.filter(name='a').exists())

    def test_duplicate_excel_row_is_skipped_without_overwrite(self):
        rows = [{
            'project_path': 'old', 'host_name': 'node', 'ip_address': '192.0.2.10',
            'port': '22', 'account_name': 'SSH', 'row': 2,
        }]
        self.assertEqual(validate_rows(rows)[0]['status'], 'duplicate')
        self.config.refresh_from_db()
        self.assertEqual(self.config.machine_name, 'machine')

    def test_new_endpoints_require_authentication(self):
        client = APIClient()
        for method, url in [
            ('get', '/api/projects/'), ('post', '/api/projects/'),
            ('patch', f'/api/projects/{self.config.project_node_id}/'),
            ('post', f'/api/file-configs/{self.config.pk}/inspect/'),
            ('post', '/api/file-configs/batch/'),
            ('get', '/api/file-configs/import-template/'),
            ('post', '/api/file-configs/import/'),
            ('post', '/api/file-configs/export/'),
        ]:
            self.assertEqual(getattr(client, method)(url).status_code, 401, url)


class InspectionParserTests(SimpleTestCase):
    @patch('dashboard.inspection.run_command')
    def test_linux_specs(self, run):
        from .inspection import inspect_machine
        run.side_effect = ['Linux', 'host\n4\n8192000\n1000000000\nDebian\nx86_64']
        self.assertEqual(inspect_machine(MagicMock())['platform_type'], 'linux')

    @patch('dashboard.inspection.run_command')
    def test_windows_specs(self, run):
        from .inspection import inspect_machine
        run.side_effect = [ValueError('uname unavailable'),
            '{"machine_name":"WIN-01","platform_type":"windows","cpu":8,"memory_bytes":17179869184,"disk_bytes":1000000000,"system":"Windows Server","arch":"64-bit"}']
        result = inspect_machine(MagicMock())
        self.assertEqual(result['platform_type'], 'windows')
        self.assertEqual(result['cpu'], 8)
        self.assertIn('-EncodedCommand', run.call_args.args[1])
