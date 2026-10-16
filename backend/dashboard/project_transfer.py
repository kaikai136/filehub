from io import BytesIO

from django.core import signing
from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from rest_framework.decorators import api_view, authentication_classes, permission_classes
from rest_framework.authentication import TokenAuthentication
from rest_framework.permissions import IsAuthenticated
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response

from .models import Account, FileConfig, Project
from .projects import validate_config
from .rbac import require_permission

IMPORT_COLUMNS = [('项目完整路径', 'project_path'), ('节点名称', 'host_name'), ('IP地址', 'ip_address'),
                  ('端口', 'port'), ('账号名称', 'account_name')]
EXPORT_INFO_COLUMNS = ['机器名称', 'CPU核数', '内存字节', '磁盘字节', '系统', '架构', '平台类型', '状态', '更新时间']
SALT = 'project-host-import'


def workbook_response(rows, filename):
    book = Workbook()
    sheet = book.active
    sheet.title = '主机配置'
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Host labels and paths must never become Excel formulas.
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    output = BytesIO()
    book.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Cache-Control'] = 'no-store'
    return response


def read_rows(upload):
    try:
        book = load_workbook(upload, read_only=True, data_only=False)
        try:
            sheet = book.active
            values = list(sheet.iter_rows(values_only=True))
            if not values or list(values[0][:len(IMPORT_COLUMNS)]) != [label for label, _ in IMPORT_COLUMNS]:
                raise ValidationError('表头不匹配，请使用下载的模板')
            rows = []
            for index, row in enumerate(values[1:], 2):
                if not any(value is not None for value in row):
                    continue
                record = {key: str(row[i] if i < len(row) and row[i] is not None else '').strip()
                          for i, (_, key) in enumerate(IMPORT_COLUMNS)}
                record['row'] = index
                rows.append(record)
            return rows
        finally:
            book.close()
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError('无法读取 Excel 文件，请上传有效的 .xlsx 文件') from error


def validate_rows(rows):
    existing = {(c.project_node.path, c.host_name, c.ip_address, c.port)
                for c in FileConfig.objects.select_related('project_node')}
    results = []
    for source in rows:
        row = dict(source)
        row['status'], row['error'] = 'ready', ''
        try:
            validate_config(row)
            row['port'] = int(row['port'])
            parts = row['project_path'].split('/')
            if not parts or any(not part.strip() or len(part.strip()) > 150 or '\\' in part for part in parts):
                raise ValidationError('项目路径无效')
            row['project_path'] = '/'.join(part.strip() for part in parts)
            accounts = list(Account.objects.filter(name=row['account_name'])[:2])
            if len(accounts) != 1:
                raise ValidationError('账号名称不存在或有重名，请先在账号管理中处理')
            if not accounts[0].is_active:
                raise ValidationError('关联账号已禁用')
            row['account_id'] = accounts[0].id
            identity = (row['project_path'], row['host_name'], row['ip_address'], row['port'])
            if identity in existing:
                row['status'] = 'duplicate'
            existing.add(identity)
        except (ValidationError, ValueError, TypeError) as error:
            row['status'], row['error'] = 'error', str(error.detail[0] if isinstance(error, ValidationError) and isinstance(error.detail, list) else error)
        results.append(row)
    return results


@api_view(['GET'])
@authentication_classes([TokenAuthentication])
@permission_classes([IsAuthenticated])
@require_permission('projects.import')
def import_template(request):
    return workbook_response([[label for label, _ in IMPORT_COLUMNS]], 'host-template.xlsx')


@api_view(['POST'])
@authentication_classes([TokenAuthentication])
@permission_classes([IsAuthenticated])
@require_permission('projects.export')
def export_configs(request):
    ids = request.data.get('ids')
    if not isinstance(ids, list) or any(not isinstance(value, int) for value in ids):
        raise ValidationError('请选择导出范围')
    rows = [[label for label, _ in IMPORT_COLUMNS] + EXPORT_INFO_COLUMNS]
    for config in FileConfig.objects.filter(id__in=ids).select_related('project_node', 'account'):
        spec = config.specifications
        rows.append([config.project_node.path, config.host_name, config.ip_address,
                     config.port, config.ssh_account.name, config.machine_name,
                     spec.get('cpu'), spec.get('memory_bytes'), spec.get('disk_bytes'), spec.get('system'),
                     spec.get('arch'), config.platform_type, config.connection_status,
                     config.updated_at.isoformat()])
    return workbook_response(rows, 'hosts.xlsx')


@api_view(['POST'])
@authentication_classes([TokenAuthentication])
@permission_classes([IsAuthenticated])
@require_permission('projects.import')
def import_configs(request):
    if 'token' not in request.data:
        upload = request.FILES.get('file')
        if upload is None or not upload.name.lower().endswith('.xlsx'):
            raise ValidationError('请选择 .xlsx 文件')
        rows = read_rows(upload)
        results = validate_rows(rows)
        valid = bool(results) and not any(row['status'] == 'error' for row in results)
        token = signing.dumps({'user': request.user.pk, 'rows': rows}, salt=SALT, compress=True) if valid else None
        return Response({'rows': results, 'token': token, 'valid': valid})
    try:
        payload = signing.loads(request.data['token'], salt=SALT, max_age=1800)
        if payload['user'] != request.user.pk:
            raise signing.BadSignature()
    except (signing.BadSignature, TypeError, KeyError):
        raise ValidationError('导入预览已失效，请重新上传')
    with transaction.atomic():
        # Revalidate against current projects, accounts and duplicate configurations.
        list(Project.objects.select_for_update().values_list('id', flat=True))
        list(Account.objects.select_for_update().values_list('id', flat=True))
        rows = validate_rows(payload['rows'])
        if any(row['status'] == 'error' for row in rows):
            return Response({'detail': '数据已变化，请重新校验', 'rows': rows}, status=400)
        created = 0
        for row in rows:
            if row['status'] == 'duplicate':
                continue
            parent = None
            for name in row['project_path'].split('/'):
                parent, _ = Project.objects.get_or_create(parent=parent, name=name)
            FileConfig.objects.create(project_node=parent, project=parent.name,
                **{key: row[key] for key in ('host_name', 'ip_address', 'port', 'account_id')},
                machine_name='', file_path='')
            created += 1
    return Response({'created': created, 'skipped': len(rows) - created})
